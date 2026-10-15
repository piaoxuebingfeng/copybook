"""拼音字帖验收 —— 从产物里量，不信生成器的自报。

版式（一方格一个音节）下要抓的错，和 v1「一行一个音节」不完全一样：

  ① 三带不等高 / 子行高不是整数像素 —— 带界错位，字母整条带地偏
  ② 贴图纵向偏移或高度 ≠ 三行子行之和 —— 同上，且**不报错**
  ③ 贴图放错格 / 示范格贴成了描红图 / 描红格贴成了示范图
     —— 按**像素内容**反查每张图是哪个音节的哪种角色（示范黑 / 描红灰）
  ④ PDF 里虚线没有把绿线严格三等分 / 竖分隔线位置漂了
  ⑤ 墨迹落带与渲染意图不一致（含声调必须进上带、降部必须进下带）

判据的教训（v1 踩过）：不能写「墨迹底部 ≈ 2/3」——`g` 的降部本来就该伸到
下带，判据错会冤枉正确的产物。这里的判据一律是「**与渲染意图一致**」。

用法：
    python3 verify_pinyin.py 字帖.xlsx --chars="bā 八 / mā 妈" [--pdf=字帖.pdf] \
        [--trace-times=full|N] [--font=路径]
"""
import hashlib
import os
import sys

import numpy as np
from PIL import Image

import pinyin_input as P
import render_pinyin as RP

ROOT = os.path.dirname(os.path.abspath(__file__))
ARG = {}
for _a in sys.argv[1:]:
    if _a.startswith("--"):
        _k, _sep, _v = _a.partition("=")
        ARG[_k] = _v if _sep else "1"

XLSX = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith("--") else None
if not XLSX:
    raise SystemExit("用法：python3 verify_pinyin.py 字帖.xlsx --chars=\"bā 八 / mā 妈\" "
                     "[--pdf=字帖.pdf] [--trace-times=full|N] [--font=路径]")
XLSX = XLSX if XLSX.startswith("/") else os.path.join(ROOT, XLSX)

SRC = ""
if ARG.get("--file"):
    _p = ARG["--file"] if ARG["--file"].startswith("/") else os.path.join(ROOT, ARG["--file"])
    SRC = open(_p, encoding="utf-8").read()
else:
    SRC = ARG.get("--chars", ARG.get("--items", ""))
ITEMS = [(i, h) for _, lst in P.parse_items(SRC) for i, h in lst]
if not ITEMS:
    raise SystemExit("✗ --chars/--file 没有解析出音节")
_bad = P.validate(ITEMS)
if _bad:
    raise SystemExit("✗ 输入里有非法音节：" + "、".join(_bad))

PDF = ARG.get("--pdf")
PDF = PDF if not PDF or PDF.startswith("/") else os.path.join(ROOT, PDF)
FONT = RP.find_font(P.needed_chars(ITEMS), ARG.get("--font"))
GEO = RP.cell_plan([i for i, _ in ITEMS], FONT)
BANDS, CELL_PX, COLS = GEO["bands"], GEO["cell_px"], GEO["cols"]
BOX_W, BOX_H, INSET_X = GEO["box_w"], GEO["box_h"], GEO["inset_x"]
BAND_PX = GEO["band_px"]

_tt = ARG.get("--trace-times", "full")
N_TRACE = COLS - 1 if _tt == "full" else max(0, min(int(_tt), COLS - 1))

# 渲染意图（验收的对照标准）：每个音节的示范图与描红图
INTENT = {}
for it, _hz in ITEMS:
    rt = RP.normalize_text(it, FONT)
    INTENT[it] = {
        "render": rt,
        "model": RP.render_cell(rt, FONT, BOX_W, BOX_H, RP.INK),
        "trace": RP.render_cell(rt, FONT, BOX_W, BOX_H, RP.TRACE),
    }

fails = []


def check(cond, ok_msg, bad_msg):
    if cond:
        print(f"  ✓ {ok_msg}")
    else:
        print(f"  ✗ {bad_msg}")
        fails.append(bad_msg)
    return cond


def pixhash(img):
    """像素内容哈希 —— 认内容不认文件名（防贴错图）。"""
    return hashlib.md5(np.asarray(img.convert("RGBA")).tobytes()).hexdigest()


# ================================================================ xlsx
print(f"=== xlsx 几何与贴图 ===  {os.path.basename(XLSX)}")
from openpyxl import load_workbook

wb = load_workbook(XLSX)
ws = wb.active


def find_bands(ws):
    """按边框模式找四线三格：行 r 顶部有边线、底部虚线 -> r 是带上沿。

    认「边框长什么样」而不是认行号 —— 行号会随排版改动而变。
    """
    out = []
    for r in range(1, ws.max_row + 1):
        bd = ws.cell(row=r, column=1).border
        nxt = ws.cell(row=r + 1, column=1).border if r + 1 <= ws.max_row else None
        if (bd.top is not None and bd.bottom is not None
                and bd.bottom.style == "dashed"
                and nxt is not None and nxt.bottom is not None):
            out.append(r)
    return [r for r in out if r + BANDS - 1 <= ws.max_row]


band_rows = find_bands(ws)
check(len(band_rows) == len(ITEMS),
      f"识别到 {len(band_rows)} 个四线三格，清单里有 {len(ITEMS)} 个音节",
      f"四线三格数量 {len(band_rows)} != 音节数 {len(ITEMS)}")

# ── ① 三带等高 ──
okh = True
for r in band_rows:
    hs = [ws.row_dimensions[r + b].height for b in range(BANDS)]
    if any(h is None or abs(h - hs[0]) > 0.01 for h in hs):
        okh = False
exp_sub = BAND_PX * 0.75                     # band_px * 72/96 = 子行高（pt）
okv = all(abs(ws.row_dimensions[r].height - exp_sub) < 0.01 for r in band_rows)
okv = all(abs(ws.row_dimensions[r].height - exp_sub) < 0.01 for r in band_rows)
check(okh and okv,
      f"全部 {len(band_rows)} 个块三带等高（{exp_sub}pt × {BANDS}）",
      "三带不等高或子行高不是整数像素量化值")

# ── ② 格子列数 ──
# 注意 openpyxl 的坑：没设过边框的单元格，border.right 是「style=None 的 Side 对象」
# 而不是 None —— 判 style 才对，判对象本身永远为真。
ncols_found = set()
for r in band_rows:
    n = 0
    while (ws.cell(row=r, column=n + 1).border.right is not None
           and ws.cell(row=r, column=n + 1).border.right.style is not None):
        n += 1
        if n > 30:
            break
    ncols_found.add(n)
check(ncols_found == {COLS},
      f"每行 {COLS} 格（格宽 {CELL_PX}px）",
      f"列数不对：{sorted(ncols_found)}（应为 {COLS}）")

# ── ③/④ 贴图：位置、尺寸、内容、角色 ──
imgs = {}
for im in ws._images:
    fr = im.anchor._from
    imgs.setdefault(fr.row, []).append((fr.col, fr.colOff, fr.rowOff,
                                        im.anchor.ext.cx, im.anchor.ext.cy, im))

EMU_PX = 9525
ok_pos = ok_role = True
roles = {}
for bi, r in enumerate(band_rows):
    it, hz = ITEMS[bi]
    got = {c: (off_c, off_r, cx, cy, im) for c, off_c, off_r, cx, cy, im in imgs.get(r - 1, [])}
    for k in range(COLS):
        if k > N_TRACE:
            if k in got:
                ok_pos = False
            continue
        if k not in got:
            ok_pos = False
            continue
        off_c, off_r, cx, cy, im = got[k]
        if off_r != 0 or abs(cy - BOX_H * EMU_PX) > EMU_PX or abs(cx - BOX_W * EMU_PX) > EMU_PX:
            ok_pos = False
        # 内容反查：示范还是描红？
        arr = np.asarray(Image.open(im.ref).convert("RGBA"))
        want_model = np.asarray(INTENT[it]["model"])
        want_trace = np.asarray(INTENT[it]["trace"])
        if arr.shape != want_model.shape:
            ok_role = False
            continue
        if np.array_equal(arr, want_model):
            roles.setdefault(bi, []).append((k, "model"))
        elif np.array_equal(arr, want_trace):
            roles.setdefault(bi, []).append((k, "trace"))
        else:
            ok_role = False
check(ok_pos,
      f"全部贴图纵向零偏移、高度 {BOX_H}px = 三行子行之和、宽 {BOX_W}px，且没有多余的图",
      "有贴图位置/尺寸不对，或超出描红范围的格子里有图")
okm = all(roles.get(bi) and roles[bi][0][1] == "model" for bi in range(len(band_rows)))
okt = all(sorted(k for k, ro in roles.get(bi, [])) == list(range(N_TRACE + 1))
          and all(ro == "trace" for k, ro in roles[bi][1:]) for bi in range(len(band_rows)))
check(ok_role and okm and okt,
      f"全部 {len(band_rows)} 格行按像素反查命中：首格示范黑 + {N_TRACE} 格描红灰",
      "有格子贴错了图（角色不符 / 顺序不对 / 内容对不上渲染意图）")

# ── ⑤ 标签行 ──
okl = True
for bi, r in enumerate(band_rows):
    v = ws.cell(row=r - 1, column=1).value
    if not v or ITEMS[bi][0] not in str(v):
        okl = False
check(okl, "每个四线格上方的标签行都写着对应音节", "有标签行不含对应音节")

if fails:
    print(f"\n✗ xlsx 验收未通过（{len(fails)} 项）")
    sys.exit(1)
print("  ✓ xlsx 通过")

# ================================================================ PDF
if not PDF or not os.path.exists(PDF):
    print("\n（未提供 --pdf，跳过 PDF 层验收）")
    print("\n✅ 拼音字帖验收通过（xlsx 层）")
    sys.exit(0)

print(f"\n=== PDF 真实像素 ===  {os.path.basename(PDF)}")
import pymupdf

DPI = 200
SCALE = DPI / 96.0
doc = pymupdf.open(PDF)


def _line_kind(green_row):
    """一条横线是实线还是虚线：数绿色段的个数。"""
    seg = green_row.astype(int)
    return int((seg[1:] & ~seg[:-1]).sum())


def page_bands(a, green):
    """把一页里的绿线切成一个个四线格。

    每带固定是 [实, 虚, 虚, 实] 四条线 —— 按这个模式切，不靠「带间距比带内
    线距大」这种假设（标签行 + 间隔行只比一条带宽窄一点，按间隙分组必错）。
    """
    H, W = green.shape
    # 阈值 50%：实线覆盖率 ~92%，虚线因为有间隙只有 ~60% —— 阈值太高会把虚线整条丢掉
    rowcov = green.sum(axis=1)
    rows = [y for y in range(H) if rowcov[y] > W * 0.5]
    lines = []
    for y in rows:
        if lines and y - lines[-1][0] <= 4:
            lines[-1] = (y, _line_kind(green[y]))          # 同一条线的抗锯齿行
        else:
            lines.append((y, _line_kind(green[y])))
    kinds = ["solid" if k < 5 else "dash" for _, k in lines]
    bands = []
    i = 0
    while i + 3 < len(lines):
        if [kinds[i], kinds[i + 1], kinds[i + 2], kinds[i + 3]] == ["solid", "dash", "dash", "solid"]:
            bands.append((lines[i][0], lines[i + 1][0], lines[i + 2][0], lines[i + 3][0]))
            i += 4
        else:
            i += 1
    return bands


all_bands = []
for pno in range(len(doc)):
    pix = doc[pno].get_pixmap(dpi=DPI)
    a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
        pix.height, pix.width, pix.n)[..., :3].astype(int)
    r_, g_, b_ = a[..., 0], a[..., 1], a[..., 2]
    green = (g_ > 180) & (g_ - r_ > 15) & (g_ - b_ > 15)
    for ys4 in page_bands(a, green):
        all_bands.append({"page": pno, "img": a, "green": green, "rows": list(ys4)})
check(len(all_bands) == len(ITEMS),
      f"{len(doc)} 页共 {len(all_bands)} 个四线格",
      f"PDF 里四线格数量 {len(all_bands)} != {len(ITEMS)}")

# ── ① 三等分 + 实虚模式 ──
oktri = True
for bd in all_bands:
    d = np.diff(bd["rows"])
    if max(d) - min(d) > 3:                       # 200dpi 下 3px
        oktri = False
check(oktri, "虚线严格三等分（≤3px@200dpi）", "有四线格的虚线没把绿线三等分")

# ── ② 竖分隔线 ──
def vlines_of(bd):
    """量这个四线格的竖分隔线位置（测出来的，不是算出来的）。"""
    gr = bd["green"]
    ys = bd["rows"]
    colcov = gr[ys[0] + 2:ys[3] - 2].sum(axis=0)
    vcols = [x for x in range(gr.shape[1]) if colcov[x] > (ys[3] - ys[0] - 4) * 0.7]
    vg = []
    for x in vcols:
        if not vg or x - vg[-1][-1] > 4:
            vg.append([x])
        else:
            vg[-1].append(x)
    return [int(np.mean(g_)) for g_ in vg]


okv = True
for bd in all_bands:
    vxs = vlines_of(bd)
    if len(vxs) != COLS + 1 or np.diff(vxs).std() > 6:
        okv = False
check(okv, f"竖分隔线每行 {COLS + 1} 根、间距均匀（格宽 ≈ {CELL_PX}px）",
      "竖分隔线数量或间距不对")

# ── ③ 墨迹落带 + 示范/描红深浅 ──
okink = okcol = True
n_cells = 0
for bi, bd in enumerate(all_bands):
    if bi >= len(ITEMS):
        break
    it, hz = ITEMS[bi]
    a = bd["img"]
    y0, y1 = bd["rows"][0], bd["rows"][3]
    vxs = vlines_of(bd)
    if len(vxs) != COLS + 1:
        okink = False
        continue
    green = bd["green"]
    ink = (a.sum(axis=2) < 700) & ~green
    for k in range(N_TRACE + 1):
        x0, x1 = vxs[k] + 4, vxs[k + 1] - 4
        cell = ink[y0 + 2:y1 - 2, x0:x1]
        if not cell.any():
            okink = False
            continue
        yy = np.flatnonzero(cell.any(axis=1))
        want = INTENT[it]["model"] if k == 0 else INTENT[it]["trace"]
        # 意图侧用 alpha>40（看得见的墨），不能用 alpha>0 —— LANCZOS 的半透明
        # 羽边在 PDF 里渲染成接近白色，两侧口径不一致就会永远差几像素
        wa = np.asarray(want)[..., 3] > 40
        wy = np.flatnonzero(wa.any(axis=1))
        # yy/wy 都是各自切片内的相对行号；切片顶 = 带上沿，意图图顶 = 贴图顶 = 同一处
        if abs(yy.min() - wy.min() * SCALE) > 5 or \
           abs(yy.max() - wy.max() * SCALE) > 5:
            okink = False
        px = a[y0 + 2:y1 - 2, x0:x1][cell]
        mean = px.mean()
        if (k == 0 and not mean < 120) or (k > 0 and not mean > 150):
            okcol = False
        n_cells += 1
check(okink,
      f"{n_cells} 个格子的墨迹位置与渲染意图一致（≤5px@200dpi，含声调进上带/降部进下带）",
      "有格子的墨迹位置与渲染意图不一致")
check(okcol, "首格墨色深（示范）、其余格墨色浅（描红）", "示范/描红的深浅不对")

if fails:
    print(f"\n✗ 拼音字帖验收未通过（{len(fails)} 项）")
    sys.exit(1)
print("\n✅ 拼音字帖验收通过（三带等高 / 贴图对齐 / 逐格内容与角色 / 严格三等分 / 落带正确）")
sys.exit(0)
